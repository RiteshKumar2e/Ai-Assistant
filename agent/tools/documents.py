"""
agent/tools/documents.py — read, chunk, retrieve, summarise and create documents.

Formats: PDF (pdfplumber → pypdf), DOCX, XLSX, PPTX, CSV, JSON, and any text or
source file. Large documents are never pasted whole into a prompt:
  * ask()        splits into overlapping chunks and scores them against the
                 question (BM25-style term weighting) — only the best few go
                 to the model, with their chunk numbers so answers can cite
  * summarize()  map-reduce: summarise chunks, then summarise the summaries
"""
from __future__ import annotations

import csv
import io
import json
import math
import re
from collections import Counter
from pathlib import Path

from agent.tools.base import ToolResult, is_secret_path, redact

CHUNK, OVERLAP = 3500, 300
_WORD = re.compile(r"[A-Za-z0-9À-￿]{2,}")
_STOP = set("the a an and or of to in on for is are was were be by with as at it this that from which what how why who "
            "does do did can could should would about into than then there their its".split())


def extract_text(p: Path, max_chars: int = 2_000_000) -> str:
    ext = p.suffix.lower()
    if ext == ".pdf":
        try:
            import pdfplumber
            with pdfplumber.open(p) as pdf:
                return "\n\n".join(f"[page {i}]\n{pg.extract_text() or ''}" for i, pg in enumerate(pdf.pages, 1))[:max_chars]
        except ImportError:
            from pypdf import PdfReader
            return "\n\n".join(f"[page {i}]\n{pg.extract_text() or ''}" for i, pg in enumerate(PdfReader(str(p)).pages, 1))[:max_chars]
    if ext == ".docx":
        import docx
        d = docx.Document(str(p))
        parts = [para.text for para in d.paragraphs]
        for t in d.tables:
            parts += [" | ".join(c.text for c in row.cells) for row in t.rows]
        return "\n".join(parts)[:max_chars]
    if ext in (".xlsx", ".xlsm"):
        import openpyxl
        wb = openpyxl.load_workbook(str(p), read_only=True, data_only=True)
        out = []
        for ws in wb.worksheets:
            out.append(f"[sheet {ws.title}]")
            for i, row in enumerate(ws.iter_rows(values_only=True)):
                if i >= 2000:
                    out.append("…"); break
                out.append(" | ".join("" if v is None else str(v) for v in row))
        return "\n".join(out)[:max_chars]
    if ext == ".pptx":
        from pptx import Presentation
        return "\n\n".join(f"[slide {i}]\n" + "\n".join(sh.text for sh in s.shapes if getattr(sh, "has_text_frame", False))
                           for i, s in enumerate(Presentation(str(p)).slides, 1))[:max_chars]
    text = p.read_text(encoding="utf-8", errors="replace")[:max_chars]
    if ext == ".json":
        try:
            return json.dumps(json.loads(text), indent=1, ensure_ascii=False)[:max_chars]
        except json.JSONDecodeError:
            return text
    return text


def chunks(text: str, size: int = CHUNK, overlap: int = OVERLAP) -> list[str]:
    if len(text) <= size:
        return [text]
    out, i = [], 0
    while i < len(text):
        end = min(len(text), i + size)
        if end < len(text):                               # break on a paragraph/sentence edge when possible
            cut = max(text.rfind("\n\n", i, end), text.rfind(". ", i, end))
            end = cut + 1 if cut > i + size // 2 else end
        out.append(text[i:end])
        if end >= len(text):
            break
        i = end - overlap
    return out


def _terms(s: str) -> list[str]:
    return [w.lower() for w in _WORD.findall(s) if w.lower() not in _STOP]


def retrieve(question: str, parts: list[str], k: int = 4) -> list[tuple[int, str]]:
    """Top-k chunks by BM25."""
    q = _terms(question)
    docs = [Counter(_terms(c)) for c in parts]
    n, avg = len(docs), (sum(sum(d.values()) for d in docs) / max(1, len(docs))) or 1
    df = Counter(t for d in docs for t in set(d))
    def bm25(d):
        L = sum(d.values())
        return sum(math.log(1 + (n - df[t] + .5) / (df[t] + .5)) * d[t] * 2.2 / (d[t] + 1.2 * (.25 + .75 * L / avg)) for t in q if t in d)
    ranked = sorted(range(n), key=lambda i: -bm25(docs[i]))[:k]
    return [(i, parts[i]) for i in sorted(ranked)]


class Documents:
    def __init__(self, fs, llm=None):
        self.fs, self.llm = fs, llm

    def _load(self, path: str) -> tuple[Path | None, str, ToolResult | None]:
        p = self.fs.resolve(path)
        if is_secret_path(p):
            return None, "", ToolResult.fail(f"{p.name} looks like a credential file — not reading it.", "denied")
        if not p.is_file():
            return None, "", ToolResult.fail(f"No such file: {p}")
        try:
            text = extract_text(p)
        except ImportError as e:
            return None, "", ToolResult.fail(f"Reading {p.suffix} needs a library that isn't installed: {e}", "unavailable")
        except Exception as e:
            return None, "", ToolResult.fail(f"Could not read {p.name}: {e}")
        if not text.strip():
            return None, "", ToolResult.fail(f"{p.name} has no extractable text (scanned/image-only?).")
        return p, redact(text), None

    def read(self, path: str, max_chars: int = 12000) -> ToolResult:
        p, text, err = self._load(path)
        if err:
            return err
        n = len(chunks(text))
        body = text[:max_chars] + (f"\n…[{len(text):,} chars total in {n} chunks — use ask_document / summarize_document]" if len(text) > max_chars else "")
        return ToolResult(True, f"{p.name} ({len(text):,} chars)\n{body}", {"path": str(p), "chunks": n})

    def ask(self, path: str, question: str) -> ToolResult:
        p, text, err = self._load(path)
        if err:
            return err
        parts = chunks(text)
        hits = retrieve(question, parts, k=4)
        ctx = "\n\n".join(f"[chunk {i + 1}/{len(parts)}]\n{c}" for i, c in hits)
        if not self.llm:
            return ToolResult(True, f"Most relevant passages of {p.name}:\n{ctx[:10000]}")
        ans = self.llm.complete(f"Answer the question using ONLY these excerpts from {p.name}. Cite chunk numbers. "
                                f"If the excerpts don't contain the answer, say so.\n\nQUESTION: {question}\n\n{ctx}")
        return ToolResult(True, ans, {"path": str(p), "chunks_used": [i + 1 for i, _ in hits]})

    def summarize(self, path: str, focus: str = "") -> ToolResult:
        p, text, err = self._load(path)
        if err:
            return err
        if not self.llm:
            return ToolResult.fail("No language model is configured to summarise with.", "unavailable")
        parts = chunks(text, 12000, 400)[:24]
        instr = f"Summarise{' with focus on: ' + focus if focus else ''}. Keep key facts, numbers, names, decisions."
        if len(parts) == 1:
            return ToolResult(True, self.llm.complete(f"{instr}\n\n{parts[0]}"), {"path": str(p)})
        partials = [self.llm.complete(f"{instr} This is part {i}/{len(parts)} of {p.name}.\n\n{c}") for i, c in enumerate(parts, 1)]
        final = self.llm.complete(f"Combine these partial summaries of {p.name} into one coherent summary. {instr}\n\n"
                                  + "\n\n".join(f"[part {i}]\n{s}" for i, s in enumerate(partials, 1)))
        note = "" if len(text) <= 24 * 12000 else f"\n\n(Summary covers the first ~{24 * 12000:,} characters of {len(text):,}.)"
        return ToolResult(True, final + note, {"path": str(p), "parts": len(parts)})

    def create(self, path: str, content: str, fmt: str = "") -> ToolResult:
        """md/txt/json/csv as text; docx via python-docx (markdown headings → Word headings)."""
        p = self.fs.resolve(path)
        fmt = (fmt or p.suffix.lstrip(".") or "md").lower()
        if fmt != "docx":
            if fmt == "csv" and isinstance(content, list):
                buf = io.StringIO(); csv.writer(buf).writerows(content); content = buf.getvalue()
            elif fmt == "json" and not isinstance(content, str):
                content = json.dumps(content, indent=2, ensure_ascii=False)
            r = self.fs.write_file(str(p), str(content))
            if r.ok:
                r.output = f"Created document {p}"
            return r
        import docx
        if p.exists():
            return ToolResult.fail(f"{p} already exists — choose another name.")
        d = docx.Document()
        for line in str(content).splitlines():
            m = re.match(r"^(#{1,4})\s+(.*)", line)
            if m:
                d.add_heading(m.group(2), level=len(m.group(1)))
            elif re.match(r"^\s*[-*]\s+", line):
                d.add_paragraph(re.sub(r"^\s*[-*]\s+", "", line), style="List Bullet")
            elif line.strip():
                d.add_paragraph(line)
        p.parent.mkdir(parents=True, exist_ok=True)
        d.save(str(p))
        return ToolResult(True, f"Created document {p}", {"path": str(p), "modified": str(p)})
